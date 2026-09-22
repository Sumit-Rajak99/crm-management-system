import os
from openpyxl import Workbook, load_workbook
from django.conf import settings


EXCEL_DIR = os.path.join(settings.BASE_DIR, "excel_data")

os.makedirs(EXCEL_DIR, exist_ok=True)


def add_data_to_excel(filename, headers, data):
    file_path = os.path.join(EXCEL_DIR, filename)

    if os.path.exists(file_path):
        workbook = load_workbook(file_path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)

    sheet.append(data)

    workbook.save(file_path)


def save_lead_to_excel(lead):

    add_data_to_excel(
        "leads.xlsx",
        [
            "ID",
            "Name",
            "Email",
            "Phone",
            "Status"
        ],
        [
            lead.id,
            lead.name,
            lead.email,
            lead.phone,
            lead.status
        ]
    )


def save_call_to_excel(call):

    add_data_to_excel(
        "calls.xlsx",
        [
            "ID",
            "Lead",
            "Employee",
            "Call Type",
            "Outcome"
        ],
        [
            call.id,
            str(call.lead),
            str(call.employee),
            call.call_type,
            call.outcome
        ]
    )


def save_followup_to_excel(followup):

    add_data_to_excel(
        "followups.xlsx",
        [
            "ID",
            "Lead",
            "Follow Up Date",
            "Status"
        ],
        [
            followup.id,
            str(followup.lead),
            followup.follow_up_date,
            followup.status
        ]
    )


def save_conversion_to_excel(conversion):

    add_data_to_excel(
        "conversions.xlsx",
        [
            "ID",
            "Lead",
            "Amount",
            "Date"
        ],
        [
            conversion.id,
            str(conversion.lead),
            conversion.amount,
            conversion.date
        ]
    )


def save_employee_to_excel(employee):

    add_data_to_excel(
        "employees.xlsx",
        [
            "ID",
            "Employee ID",
            "Full Name",
            "Department",
            "Designation",
            "Role"
        ],
        [
            employee.id,
            employee.employee_id,
            employee.full_name,
            employee.department,
            employee.designation,
            employee.role
        ]
    )